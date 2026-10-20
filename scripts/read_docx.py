import docx
import os

doc_path = r"d:\ProgramFiles\GitHub-Firebase-Antigravity\2026-07 Paridad Arancel Dialitico\directivas\vision\Trabajo_UNR_Demartis Pablo Nicolás_PARIDAD TÉCNICA ARANCEL DIALÍTICO.docx"
output_path = r"d:\ProgramFiles\GitHub-Firebase-Antigravity\2026-07 Paridad Arancel Dialitico\.tmp\paper_extracted.txt"

os.makedirs(os.path.dirname(output_path), exist_ok=True)

try:
    doc = docx.Document(doc_path)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    
    # Also get tables text if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text for cell in row.cells]
            fullText.append(" | ".join(row_text))

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(fullText))
    
    print(f"Extracted {len(fullText)} paragraphs/tables lines successfully to {output_path}")
except Exception as e:
    print(f"Error: {e}")
